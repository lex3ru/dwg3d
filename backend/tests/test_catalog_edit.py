"""Catalog editing through the API, uploaded models, YAML export, reports."""
import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from dwg3d import catalog, db, store, viewdata
from conftest import add_drawing
from dwg3d.models import AuditLog, Build, Project


@pytest.fixture
def api(pg, tmp_path, monkeypatch):
    monkeypatch.setenv("DWG3D_DATA", str(tmp_path / "data"))
    from dwg3d import api as module
    return TestClient(module.app)


def _glb(n=64) -> bytes:
    return b"glTF" + (2).to_bytes(4, "little") + (n).to_bytes(4, "little") + b"\0" * (n - 12)


def test_create_update_deactivate_and_audit(api):
    iid = f"test.{uuid.uuid4().hex[:6]}"
    item = {"id": iid, "cat": "diagnostics", "name": "МРТ 3 Тл", "w": 2.4, "d": 5, "h": 2.4, "kind": "mri",
            "article": "MR-3T", "supplier": "ООО «Мед»", "price": 150000000, "currency": "RUB"}
    r = api.post("/api/catalog/items", json=item)
    assert r.status_code == 201 and r.json()["price"] == 150000000.0 and r.json()["cat"] == "diagnostics"
    assert api.post("/api/catalog/items", json=item).status_code == 400                 # id taken
    assert api.post("/api/catalog/items", json={**item, "id": "Bad Id"}).status_code == 400
    assert api.post("/api/catalog/items", json={**item, "id": iid + "x", "cat": "nope"}).status_code == 400
    assert api.post("/api/catalog/items", json={**item, "id": iid + "y", "w": -1}).status_code == 400

    r = api.patch(f"/api/catalog/items/{iid}", json={"price": 149000000.5, "supplier": None, "cat": "engineering"})
    assert r.status_code == 200
    got = r.json()
    assert got["price"] == 149000000.5 and "supplier" not in got and got["cat"] == "engineering"
    assert got["name"] == "МРТ 3 Тл" and got["article"] == "MR-3T"                      # others kept
    assert api.patch(f"/api/catalog/items/{iid}", json={"h": 0}).status_code == 400
    assert api.patch("/api/catalog/items/nope", json={"price": 1}).status_code == 404
    got = api.patch(f"/api/catalog/items/{iid}", json={"active": False}).json()
    assert got["active"] is False
    listed = {i["id"]: i for i in api.get("/api/catalog").json()["items"]}
    assert listed[iid]["active"] is False
    with db.session() as s:
        acts = [a.action for a in s.scalars(select(AuditLog).where(AuditLog.entity == "catalog_item",
                                                                    AuditLog.entity_id == iid).order_by(AuditLog.id))]
    assert acts == ["create", "update", "update"]


def test_model_upload_and_serving(api):
    iid = f"test.{uuid.uuid4().hex[:6]}"
    api.post("/api/catalog/items", json={"id": iid, "cat": "storage", "name": "Шкаф", "w": 1, "d": 0.5, "h": 2,
                                         "kind": "cabinet"})
    assert api.post(f"/api/catalog/items/{iid}/model", files={"file": ("x.glb", b"not a glb")}).status_code == 400
    r = api.post(f"/api/catalog/items/{iid}/model", files={"file": ("x.glb", _glb())})
    assert r.status_code == 200
    url = r.json()["modelUrl"]
    assert r.json()["model"].startswith("upload/") and url.startswith("/api/catalog/models/")
    got = api.get(url)
    assert got.status_code == 200 and got.content == _glb() and "immutable" in got.headers["cache-control"]
    assert api.get("/api/catalog/models/..%2F..%2Fjob.json").status_code == 404
    assert api.get("/api/catalog/models/missing.glb").status_code == 404
    # bundled models keep their public path
    sofa = next(i for i in api.get("/api/catalog").json()["items"] if i["id"] == "rec.sofa")
    assert sofa["modelUrl"] == "/catalog/kenney/loungeSofa.glb"


def test_categories(api):
    cid = f"cat{uuid.uuid4().hex[:6]}"
    r = api.post("/api/catalog/categories", json={"id": cid, "name": "Реабилитация", "color": "#12b886"})
    assert r.status_code == 201
    assert api.post("/api/catalog/categories", json={"id": cid, "name": "x"}).status_code == 400
    assert api.patch(f"/api/catalog/categories/{cid}", json={"name": "ЛФК", "color": "bad"}).status_code == 400
    assert api.patch(f"/api/catalog/categories/{cid}", json={"name": "ЛФК"}).json()["name"] == "ЛФК"
    assert any(c["id"] == cid for c in api.get("/api/catalog").json()["categories"])


def test_dump_round_trip(pg, tmp_path):
    catalog.load()
    before = catalog.as_json()
    path = tmp_path / "cat.yaml"
    assert catalog.dump(path)["items"] == len(before["items"])
    catalog.load(path)
    after = catalog.as_json()
    strip = lambda d: {i["id"]: {k: v for k, v in i.items() if k != "modelUrl"} for i in d["items"]}  # noqa: E731
    assert strip(after) == strip(before)


# ---------------------------------------------------------------- reports ---
def _project_with(items) -> str:
    """A project with two floors and the given (floor, item) placements."""
    pid = uuid.uuid4().hex[:12]
    floors = [{"id": "c0", "space": "Model", "owner": None, "bbox": [0, 0, 30, 20], "offset": [0, 0], "name": "1 этаж", "level": 1},
              {"id": "c1", "space": "Model", "owner": None, "bbox": [40, 0, 70, 20], "offset": [0, 0], "name": "2 этаж", "level": 2}]
    with db.session() as s:
        add_drawing(s, pid, project_name=f"Клиника {pid[:4]}", current_build=1)
        s.add(Build(drawing_id=pid, n=1, result={"floors": floors}, config={}, out_dir="out"))
        store.sync_floors(s, pid, {"floors": floors})
        viewdata.apply(s, pid, "placement", [{"op": "create", "floor": f, "item": it, "at": [1, 1]} for f, it in items])
    return pid


def test_project_specification(api):
    catalog.update_item("doc.couch", {"price": 25000, "article": "KS-1", "supplier": "Поставщик"})
    catalog.update_item("doc.chair", {"price": None})
    pid = _project_with([("c0", "doc.couch"), ("c0", "doc.couch"), ("c1", "doc.couch"), ("c1", "doc.chair")])
    rep = api.get("/api/reports/equipment", params={"project": pid}).json()
    assert rep["project"]["name"].startswith("Клиника") and rep["columns"] == ["1 этаж", "2 этаж"]
    rows = {r["item"]: r for r in rep["rows"]}
    assert rows["doc.couch"]["qty"] == 3 and rows["doc.couch"]["by"] == {"1 этаж": 2, "2 этаж": 1}
    assert rows["doc.couch"]["sum"] == 75000 and rows["doc.couch"]["article"] == "KS-1"
    assert rows["doc.chair"]["price"] is None and rows["doc.chair"]["sum"] is None
    assert rep["totals"] == {"RUB": 75000} and rep["qty"] == 4 and rep["unpriced"] == 1
    assert api.get("/api/reports/equipment", params={"project": "nope"}).status_code == 404

    r = api.get("/api/reports/equipment", params={"project": pid, "format": "xlsx"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    cells = [[c.value for c in row] for row in ws.iter_rows()]
    header = next(row for row in cells if row and row[0] == "№")
    assert header[6:8] == ["1 этаж", "2 этаж"] and header[8:] == ["Кол-во", "Цена", "Сумма", "Валюта"]
    couch = next(row for row in cells if row[2] and "Кушетка" in row[2])
    assert couch[6:8] == [2, 1] and couch[8].startswith("=SUM(") and couch[9] == 25000 and couch[10].startswith("=")
    assert any(row[2] == "Итого" for row in cells) and any(row[2] and "Без цены" in str(row[2]) for row in cells)


def test_summary_over_projects(api):
    a = _project_with([("c0", "diag.mri")])
    b = _project_with([("c0", "diag.mri"), ("c1", "diag.mri"), ("c1", "diag.ct")])
    rep = api.get("/api/reports/equipment").json()
    ids = [p["id"] for p in rep["projects"]]
    assert a in ids and b in ids and ids.index(b) < ids.index(a)       # busiest project first
    c = _project_with([("c0", "diag.ct")])
    with db.session() as s:
        s.get(Project, c).name = s.get(Project, a).name       # same name as another building
    names = {p["id"]: p["name"] for p in api.get("/api/reports/equipment").json()["projects"]}
    assert names[a].endswith(a[:6]) and names[c].endswith(c[:6]) and not names[b].endswith(b[:6])
    mri = next(r for r in rep["rows"] if r["item"] == "diag.mri")
    assert mri["by"][a] == 1 and mri["by"][b] == 2
    r = api.get("/api/reports/equipment", params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value.startswith("Сводная")
    assert api.get("/api/reports/equipment", params={"format": "pdf"}).status_code == 400
