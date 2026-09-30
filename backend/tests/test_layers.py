"""Layers of a project, their objects, rooms that keep their ids, the room report."""
import io
import os
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from make_sample import make


@pytest.fixture(scope="module")
def c(tmp_path_factory, pg):
    os.environ["DWG3D_DATA"] = str(tmp_path_factory.mktemp("data"))
    from dwg3d import api
    return TestClient(api.app)


@pytest.fixture(scope="module")
def built(c, tmp_path_factory):
    """A project with one built drawing: (project id, drawing JSON)."""
    sample = make(tmp_path_factory.mktemp("dxf") / "sample.dxf")
    pid = c.post("/api/projects", json={"name": "Слои"}).json()["id"]
    did = c.post(f"/api/projects/{pid}/drawings", files={"file": ("plan.dxf", sample.read_bytes())}).json()["id"]
    return pid, _wait(c, did)


def _wait(c, did, timeout=120):
    t = time.time()
    while time.time() - t < timeout:
        d = c.get(f"/api/drawings/{did}").json()
        if d["status"] in ("done", "error"):
            assert d["status"] == "done", d.get("error")
            return d
        time.sleep(0.2)
    raise TimeoutError


def test_layer_kinds_and_crud(c, built):
    pid, d = built
    kinds = {k["kind"]: k for k in c.get("/api/layer-kinds").json()}
    assert {"furniture", "cctv"} <= set(kinds) and kinds["cctv"]["categories"] == ["security"]
    assert c.post(f"/api/projects/{pid}/layers", json={"kind": "nope"}).status_code == 400
    lay = c.post(f"/api/projects/{pid}/layers", json={"kind": "furniture", "name": "Мебель — вариант Б"})
    assert lay.status_code == 201
    lay = lay.json()
    assert lay["name"] == "Мебель — вариант Б" and lay["visible"] and not lay["locked"] and lay["count"] == 0
    r = c.patch(f"/api/projects/{pid}/layers/{lay['id']}", json={"visible": False, "version": lay["version"]})
    assert r.status_code == 200 and r.json()["visible"] is False and r.json()["version"] == lay["version"] + 1
    assert c.patch(f"/api/projects/{pid}/layers/{lay['id']}", json={"name": "x", "version": lay["version"]}).status_code == 409
    assert c.patch(f"/api/projects/{pid}/layers/{lay['id']}", json={"visible": "yes"}).status_code == 400
    assert any(x["id"] == lay["id"] for x in c.get(f"/api/projects/{pid}/layers").json())
    assert c.delete(f"/api/projects/{pid}/layers/{lay['id']}").status_code == 204
    assert c.patch(f"/api/projects/{pid}/layers/{lay['id']}", json={"name": "y"}).status_code == 404


def test_objects_in_layers(c, built):
    pid, d = built
    did, fid = d["id"], d["result"]["floors"][0]["id"]
    base = f"/api/drawings/{did}/objects"
    couch = c.post(base, json={"id": "couch", "floor": fid, "item": "doc.couch", "at": [2, 2]}).json()
    cam = c.post(base, json={"id": "cam", "floor": fid, "item": "sec.dome", "at": [3, 3]}).json()
    layers = {x["id"]: x for x in c.get(f"/api/projects/{pid}/layers").json()}
    assert layers[couch["layer"]]["kind"] == "furniture" and layers[cam["layer"]]["kind"] == "cctv"
    cctv = cam["layer"]
    # the CCTV layer takes only CCTV equipment
    r = c.post(base, json={"id": "c2", "floor": fid, "item": "doc.couch", "at": [2, 2], "layer": cctv})
    assert r.status_code == 400 and "не принимает" in r.json()["detail"]
    assert c.post(base, json={"floor": fid, "item": "doc.couch", "at": [2, 2], "layer": "nope"}).status_code == 400
    # polylines are not furniture
    r = c.post(base, json={"floor": fid, "geom": "polyline", "points": [[0, 0], [1, 1]], "layer": couch["layer"]})
    assert r.status_code == 400
    # a second furniture layer: objects move between layers of a kind that accepts them
    b = c.post(f"/api/projects/{pid}/layers", json={"kind": "furniture", "name": "Вариант Б"}).json()
    moved = c.patch(f"{base}/couch", json={"layer": b["id"], "version": couch["version"]}).json()
    assert moved["layer"] == b["id"] and moved["version"] == couch["version"] + 1
    items = c.get(f"{base}?layer={b['id']}").json()["items"]
    assert [x["id"] for x in items] == ["couch"]
    # a locked layer: its objects cannot be changed or deleted, nor can the layer
    c.patch(f"/api/projects/{pid}/layers/{b['id']}", json={"locked": True})
    r = c.patch(f"{base}/couch", json={"at": [5, 5], "version": moved["version"]})
    assert r.status_code == 400 and "заблокирован" in r.json()["detail"]
    assert c.delete(f"{base}/couch?version={moved['version']}").status_code == 400
    assert c.delete(f"/api/projects/{pid}/layers/{b['id']}").status_code == 409
    c.patch(f"/api/projects/{pid}/layers/{b['id']}", json={"locked": False})
    # deleting a layer deletes its objects
    assert c.delete(f"/api/projects/{pid}/layers/{b['id']}").status_code == 204
    assert "couch" not in {x["id"] for x in c.get(base).json()["items"]}
    log = c.get(f"/api/projects/{pid}/audit").json()
    assert {("layer", "create"), ("layer", "update"), ("layer", "delete"), ("object", "create")} <= \
        {(x["entity"], x["action"]) for x in log}


def test_rooms_keep_ids_and_objects_know_their_room(c, built):
    pid, d = built
    did = d["id"]
    rooms = c.get(f"/api/drawings/{did}/rooms").json()
    assert rooms and all(r["idx"] is not None and r["floor"] for r in rooms)
    fl = d["result"]["floors"][0]
    meta = [r for r in fl["rooms"] if r.get("id")]
    assert meta and {r["id"] for r in meta} <= {r["id"] for r in rooms}
    # a point inside a room (its label position in plan coordinates, from the plan)
    plan = c.get(f"/api/drawings/{did}/files/{fl['plan']}").json()
    k, room = max(enumerate(plan["rooms"]), key=lambda kr: kr[1]["area"])
    rid = next(r["id"] for r in rooms if r["floor"] == fl["id"] and r["idx"] == k)
    obj = c.post(f"/api/drawings/{did}/objects", json={"id": "inroom", "floor": fl["id"], "item": "doc.desk",
                                                       "at": room["label"]}).json()
    assert obj["room"] == rid
    far = c.post(f"/api/drawings/{did}/objects", json={"id": "far", "floor": fl["id"], "item": "doc.desk",
                                                       "at": [-500, -500]}).json()
    assert far["room"] is None
    # a rebuild keeps room ids and the object's room
    c.post(f"/api/drawings/{did}/rebuild", json={"geometry": {"snap_tolerance": 0.011}})
    d2 = _wait(c, did)
    assert d2["build"] == d["build"] + 1
    rooms2 = c.get(f"/api/drawings/{did}/rooms").json()
    assert {r["id"] for r in rooms2 if r["idx"] is not None} == {r["id"] for r in rooms}
    objs = {x["id"]: x for x in c.get(f"/api/drawings/{did}/objects").json()["items"]}
    assert objs["inroom"]["room"] == rid

    # the specification room by room
    rep = c.get(f"/api/reports/equipment?project={pid}&by=room").json()
    names = {(g["room"], x["item"]) for g in rep["rooms"] for x in g["items"]}
    assert (room["name"] or "без названия", "doc.desk") in names and ("вне помещений", "doc.desk") in names
    assert c.get("/api/reports/equipment?by=room").status_code == 400
    x = c.get(f"/api/reports/equipment?project={pid}&by=room&format=xlsx")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    assert "по помещениям" in ws["A1"].value
    cctv = next(x for x in c.get(f"/api/projects/{pid}/layers").json() if x["kind"] == "cctv")
    rep = c.get(f"/api/reports/equipment?project={pid}&layer={cctv['id']}").json()
    assert {r["item"] for r in rep["rows"]} == {"sec.dome"} and rep["layer"]["name"] == cctv["name"]


def test_rooms_backfill(c, built, pg):
    """Drawings built before rooms had rows get them from the stored plans."""
    from sqlalchemy import delete

    from dwg3d import db, rooms
    from dwg3d.models import Build, Room

    pid, d = built
    did = d["id"]
    with db.session() as s:
        b = s.get(Build, (did, c.get(f"/api/drawings/{did}").json()["build"]))
        ids_before = {r["id"] for f in b.result["floors"] for r in f["rooms"] if r.get("id")}
        s.execute(delete(Room).where(Room.drawing_id == did))
        for f in b.result["floors"]:
            for r in f["rooms"]:
                r.pop("id", None)
    with db.session() as s:
        n = rooms.backfill(s, did)
    assert n == len(ids_before)
    d2 = c.get(f"/api/drawings/{did}").json()
    assert all(r.get("id") for f in d2["result"]["floors"] for r in f["rooms"])
