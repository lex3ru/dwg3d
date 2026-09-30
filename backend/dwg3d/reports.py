"""Equipment specification: placed items counted per catalog position.

  equipment(project_id)  one project, with the count per floor (of each drawing)
  equipment(None)        all projects, with the count per project
  by_room(project_id)    one project, room by room (for fitting out and procurement)
  (optionally of one layer)

Items whose floor is not in the current build count under "без этажа".
When a project has several drawings, floor columns carry the drawing name.
Sums are per currency; items without a price are listed and counted but
left out of the sums. `to_xlsx()` writes the same table with formulas
(quantity × price, totals), so edits in Excel recalculate.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .models import CatalogCategory, CatalogItem, Drawing, Floor, Layer, LayerObject, Project, Room

NO_FLOOR = "без этажа"


def _project_name(p: Project) -> str:
    return p.name


def _item_row(it: CatalogItem | None, iid: str, cats: dict) -> dict:
    cat = cats.get(it.category_id) if it else None
    return {"item": iid, "name": it.name if it else f"Неизвестный предмет ({iid})",
            "category": cat.name if cat else "", "category_sort": cat.sort if cat else 1e9,
            "sort": it.sort if it else 1e9, "article": (it.article if it else None) or "",
            "supplier": (it.supplier if it else None) or "",
            "price": float(it.price) if it and it.price is not None else None,
            "currency": it.currency if it and it.price is not None else None, "qty": 0}


def _placed(q, project_id: str | None, layer_id: str | None, me=None, layer_ids: set[str] | None = None):
    """Catalog objects of live projects and drawings (one project / layer;
    only projects `me` may open and layers in `layer_ids` when given)."""
    q = (q.join(Drawing, Drawing.id == LayerObject.drawing_id)
         .join(Project, Project.id == Drawing.project_id)
         .where(Project.deleted_at.is_(None), Drawing.deleted_at.is_(None), LayerObject.item_id.is_not(None)))
    if me is not None:
        from . import access
        q = access.visible(q, me)
        if me.share is not None and me.share["drawings"] is not None:
            q = q.where(Drawing.id.in_(me.share["drawings"]))
    if layer_ids is not None:
        q = q.where(LayerObject.layer_id.in_(layer_ids))
    if project_id is not None:
        q = q.where(Drawing.project_id == project_id)
    if layer_id is not None:
        q = q.where(LayerObject.layer_id == layer_id)
    return q


def equipment(s: Session, project_id: str | None = None, layer_id: str | None = None, me=None,
              layer_ids: set[str] | None = None) -> dict:
    q = _placed(select(Drawing.project_id, LayerObject.item_id, Floor.name, Floor.candidate, Floor.level,
                       Drawing.id, func.count())
                .select_from(LayerObject)
                .outerjoin(Floor, Floor.id == LayerObject.floor_id), project_id, layer_id, me, layer_ids)
    q = q.group_by(Drawing.project_id, LayerObject.item_id, Floor.name, Floor.candidate, Floor.level, Drawing.id)
    counts = s.execute(q).all()
    drawings = {d.id: d for d in s.scalars(select(Drawing).where(Drawing.id.in_({c[5] for c in counts})))}
    many = len(drawings) > 1
    items = {i.id: i for i in s.scalars(select(CatalogItem).where(CatalogItem.id.in_({c[1] for c in counts})))}
    cats = {c.id: c for c in s.scalars(select(CatalogCategory))}
    projects = {p.id: p for p in s.scalars(select(Project).where(Project.id.in_({c[0] for c in counts} |
                                                                                 ({project_id} if project_id else set()))))}

    rows: dict[str, dict] = {}
    floor_order: dict[str, tuple] = {}
    for pid, iid, fname, cand, level, did, n in counts:
        r = rows.get(iid)
        if r is None:
            r = rows[iid] = {**_item_row(items.get(iid), iid, cats), "by": defaultdict(int)}
        r["qty"] += n
        if project_id is not None:
            label = (fname or cand) if cand else NO_FLOOR
            dname = drawings[did].name or drawings[did].filename.rsplit(".", 1)[0]
            if many:
                label = f"{dname} · {label}"
            floor_order.setdefault(label, (dname if many else "", 0 if cand else 1,
                                           level if level is not None else 0, label))
            r["by"][label] += n
        else:
            r["by"][pid] += n

    out_rows = sorted(rows.values(), key=lambda r: (r["category_sort"], r["sort"], r["name"]))
    totals: dict[str, float] = defaultdict(float)
    for r in out_rows:
        r["sum"] = round(r["price"] * r["qty"], 2) if r["price"] is not None else None
        if r["sum"] is not None:
            totals[r["currency"]] += r["sum"]
        r["by"] = dict(r["by"])
        del r["category_sort"], r["sort"]
    if project_id is not None:
        p = projects.get(project_id)
        columns = [label for label in sorted(floor_order, key=lambda k: floor_order[k])]
        head = {"project": {"id": project_id, "name": _project_name(p) if p else project_id}}
        if layer_id is not None:
            layer = s.get(Layer, layer_id)
            head["layer"] = {"id": layer_id, "name": layer.name if layer else layer_id}
    else:
        # projects with most items first
        per = defaultdict(int)
        for r in out_rows:
            for pid, n in r["by"].items():
                per[pid] += n
        order = sorted(per, key=lambda pid: (-per[pid], _project_name(projects[pid])))
        columns = order
        names = [_project_name(projects[pid]) for pid in order]
        # buildings uploaded from the same drawing share a name: tell the columns apart
        head = {"projects": [{"id": pid, "name": f"{n} · {pid[:6]}" if names.count(n) > 1 else n}
                             for pid, n in zip(order, names, strict=True)]}
    return {**head, "columns": columns, "rows": out_rows,
            "totals": {k: round(v, 2) for k, v in totals.items()},
            "qty": sum(r["qty"] for r in out_rows),
            "unpriced": sum(1 for r in out_rows if r["price"] is None),
            "generated": datetime.now().astimezone().isoformat(timespec="minutes")}


NO_ROOM = "вне помещений"


def by_room(s: Session, project_id: str, layer_id: str | None = None, me=None,
            layer_ids: set[str] | None = None) -> dict:
    """Catalog objects of a project room by room: [{drawing, floor, room, area, items, totals}]."""
    q = _placed(select(Drawing.id, Floor.name, Floor.candidate, Floor.level, Room.id, Room.name, Room.area,
                       LayerObject.item_id, func.count())
                .select_from(LayerObject)
                .outerjoin(Floor, Floor.id == LayerObject.floor_id)
                .outerjoin(Room, Room.id == LayerObject.room_id), project_id, layer_id, me, layer_ids)
    q = q.group_by(Drawing.id, Floor.name, Floor.candidate, Floor.level, Room.id, Room.name, Room.area,
                   LayerObject.item_id)
    counts = s.execute(q).all()
    items = {i.id: i for i in s.scalars(select(CatalogItem).where(CatalogItem.id.in_({c[7] for c in counts})))}
    cats = {c.id: c for c in s.scalars(select(CatalogCategory))}
    drawings = {d.id: d for d in s.scalars(select(Drawing).where(Drawing.project_id == project_id))}
    many = len({c[0] for c in counts}) > 1
    groups: dict[tuple, dict] = {}
    for did, fname, cand, level, rid, rname, area, iid, n in counts:
        d = drawings.get(did)
        dname = (d.name or d.filename.rsplit(".", 1)[0]) if d else did
        floor = (fname or cand) if cand else NO_FLOOR
        key = (dname if many else "", 0 if cand else 1, level if level is not None else 0, floor,
               rid is None, rname or "", rid or 0)
        g = groups.get(key)
        if g is None:
            g = groups[key] = {"drawing": dname, "floor": floor, "room_id": rid,
                               "room": (rname or "без названия") if rid else NO_ROOM,
                               "area": round(area, 2) if rid and area else None, "items": {}}
        r = g["items"].get(iid)
        if r is None:
            r = g["items"][iid] = _item_row(items.get(iid), iid, cats)
        r["qty"] += n
    out, totals = [], defaultdict(float)
    for key in sorted(groups):
        g = groups[key]
        rows = sorted(g["items"].values(), key=lambda r: (r["category_sort"], r["sort"], r["name"]))
        sums = defaultdict(float)
        for r in rows:
            r["sum"] = round(r["price"] * r["qty"], 2) if r["price"] is not None else None
            if r["sum"] is not None:
                sums[r["currency"]] += r["sum"]
                totals[r["currency"]] += r["sum"]
            del r["category_sort"], r["sort"]
        out.append({**g, "items": rows, "qty": sum(r["qty"] for r in rows),
                    "totals": {k: round(v, 2) for k, v in sums.items()}})
    p = s.get(Project, project_id)
    head = {"project": {"id": project_id, "name": p.name if p else project_id}}
    if layer_id is not None:
        layer = s.get(Layer, layer_id)
        head["layer"] = {"id": layer_id, "name": layer.name if layer else layer_id}
    return {**head, "many_drawings": many, "rooms": out, "qty": sum(g["qty"] for g in out),
            "totals": {k: round(v, 2) for k, v in totals.items()},
            "generated": datetime.now().astimezone().isoformat(timespec="minutes")}


def rooms_to_xlsx(rep: dict) -> bytes:
    """The room-by-room specification: a block per room with a subtotal."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "По помещениям"
    title = f"Спецификация по помещениям: {rep['project']['name']}"
    if rep.get("layer"):
        title += f" — {rep['layer']['name']}"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Сформирована {rep['generated'][:16].replace('T', ' ')}"])
    ws.append([])
    ws.append(["№", "Наименование", "Категория", "Артикул", "Кол-во", "Цена", "Сумма", "Валюта"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E9ECEF")
    sum_cells = []
    for g in rep["rooms"]:
        where = " · ".join(x for x in ((g["drawing"] if rep.get("many_drawings") else ""), g["floor"], g["room"]) if x)
        if g.get("area"):
            where += f" ({g['area']:.1f} м²)"
        ws.append([where])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        first = ws.max_row + 1
        for k, r in enumerate(g["items"], start=1):
            row = ws.max_row + 1
            total = f"=E{row}*F{row}" if r["price"] is not None else None
            ws.append([k, r["name"], r["category"], r["article"], r["qty"], r["price"], total, r["currency"] or ""])
        last = ws.max_row
        ws.append(["", "Итого по помещению", "", "", f"=SUM(E{first}:E{last})", None,
                   f"=SUM(G{first}:G{last})" if len(g["totals"]) <= 1 else None, next(iter(g["totals"]), "")])
        ws.cell(ws.max_row, 2).font = Font(italic=True)
        sum_cells.append(ws.max_row)
        ws.append([])
    if sum_cells:
        ws.append(["", "Итого по проекту", "", "", "=" + "+".join(f"E{r}" for r in sum_cells), None,
                   ("=" + "+".join(f"G{r}" for r in sum_cells)) if len(rep["totals"]) <= 1 else None,
                   next(iter(rep["totals"]), "")])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (5, 44, 24, 16, 9, 12, 14, 8), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for c in row:
            if c.column in (6, 7):
                c.number_format = "#,##0.00"
            if c.column == 2:
                c.alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_xlsx(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    single = "project" in rep
    ws.title = "Спецификация" if single else "Сводная"
    col_names = [p["name"] for p in rep["projects"]] if not single else rep["columns"]
    title = (f"Спецификация оборудования: {rep['project']['name']}" if single
             else f"Сводная спецификация оборудования по проектам ({len(rep['projects'])})")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Сформирована {rep['generated'][:16].replace('T', ' ')}"])
    ws.append([])
    fixed = ["№", "Категория", "Наименование", "Артикул", "Поставщик", "Ед."]
    head = fixed + col_names + ["Кол-во", "Цена", "Сумма", "Валюта"]
    ws.append(head)
    hr = ws.max_row
    for c in ws[hr]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E9ECEF")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    first = hr + 1
    nb = len(col_names)
    qty_col = len(fixed) + nb + 1
    L = get_column_letter
    keys = rep["columns"]
    for k, r in enumerate(rep["rows"], start=1):
        row = first + k - 1
        by = [r["by"].get(key) or None for key in keys]
        qty = f"=SUM({L(len(fixed) + 1)}{row}:{L(len(fixed) + nb)}{row})" if nb else r["qty"]
        price = r["price"]
        total = f"={L(qty_col)}{row}*{L(qty_col + 1)}{row}" if price is not None else None
        ws.append([k, r["category"], r["name"], r["article"], r["supplier"], "шт."] + by +
                  [qty, price, total, r["currency"] or ""])
    last = ws.max_row
    if rep["rows"]:
        ws.append([])
        ws.append(["", "", "Итого", "", "", ""] + [f"=SUM({L(len(fixed) + 1 + i)}{first}:{L(len(fixed) + 1 + i)}{last})"
                                                   for i in range(nb)] +
                  [f"=SUM({L(qty_col)}{first}:{L(qty_col)}{last})", None,
                   f"=SUM({L(qty_col + 2)}{first}:{L(qty_col + 2)}{last})" if len(rep["totals"]) <= 1 else None,
                   next(iter(rep["totals"]), "")])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        if len(rep["totals"]) > 1:
            for cur, v in rep["totals"].items():
                ws.append(["", "", f"Итого, {cur}"] + [None] * (qty_col + 1 - 3) + [v, cur])
        if rep["unpriced"]:
            ws.append(["", "", f"Без цены: {rep['unpriced']} поз. (не вошли в сумму)"])
    widths = [5, 24, 44, 16, 22, 6] + [10] * nb + [9, 12, 14, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for c in row:
            if c.column in (qty_col + 1, qty_col + 2):
                c.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=first, column=4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

