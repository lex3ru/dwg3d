# /// script
# requires-python = ">=3.12"
# dependencies = ["httpx", "openpyxl"]
# ///
"""Trying to get at what is not one's own (docs/ACCEPTANCE.md, 12.4), on what seed.py made:

    uv run tools/acceptance/probe.py

As a user who sees one building (load-u50) and as an administrator of another building (load-u01):
read and change the data of a building that is not theirs through the gateway — its tree, plans, rooms,
objects, people, journal, variants, reports, rights; change what they may only see; act as the system
administrator. Then with no session, a forged cookie, a forged token; and straight at the services,
past the gateway, from inside the network: no token, a token signed by someone else, `alg: none`.
Every attempt must be refused (401, 403, 404) or come back empty. This is the automated part of the
check; the criterion asks also for a person from outside to try. Exit code 1 when anything gets through.
"""
from __future__ import annotations

import asyncio
import base64
import io
import json
import subprocess
import sys
import time
from pathlib import Path

import httpx

sys.path.insert(0, str(Path(__file__).parent))
from common import BASE, load_state, sign_in  # noqa: E402

RESULTS: list[tuple[bool, str, str]] = []


def empty(r: httpx.Response) -> bool:
    """Refused, or answered with nothing in it: no items, no floors, a cabinet with no rooms, a sheet of headers."""
    if r.status_code in (401, 403, 404):
        return True
    if r.status_code != 200:
        return False
    kind = r.headers.get("content-type", "")
    if "json" in kind:
        body = r.json()
        if isinstance(body, dict) and "items" in body:
            return not body["items"]
        if isinstance(body, dict) and "floors" in body:
            return not body["floors"]
        if isinstance(body, dict) and "rooms" in body:
            return not any(body.get(k) for k in ("rooms", "mine", "attention"))
    if "spreadsheet" in kind:
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(r.content), read_only=True).active.max_row <= 1
    return False


def refused(r: httpx.Response) -> bool:
    """A change refused: no right, or what it refers to is not there for the asker (a 422 «не найден» is the same
    answer as for something that does not exist)."""
    return r.status_code in (401, 403, 404, 405) or (r.status_code == 422 and "не найден" in r.text)


async def attempt(c: httpx.AsyncClient, who: str, what: str, method: str, path: str, **kw) -> None:
    r = await c.request(method, path, **kw)
    ok = empty(r) if method == "GET" else refused(r)
    RESULTS.append((ok, who, f"{what}: {method} {path.split('?')[0]} → {r.status_code}"))


async def as_outsider(c: httpx.AsyncClient, who: str, other: str, room: str, floor: str, asset: dict) -> None:
    a = lambda *x, **kw: attempt(c, who, *x, **kw)  # noqa: E731
    await a("чужое здание", "GET", f"/api/buildings/v1/buildings/{other}")
    await a("дерево чужого здания", "GET", f"/api/buildings/v1/buildings/{other}/tree")
    await a("модель чужого здания", "GET", f"/api/buildings/v1/buildings/{other}/model")
    await a("план чужого этажа", "GET", f"/api/buildings/v1/floors/{floor}/plan")
    await a("помещения чужого здания", "GET", f"/api/buildings/v1/rooms?building_id={other}")
    await a("чужое помещение", "GET", f"/api/buildings/v1/rooms/{room}")
    await a("объекты чужого здания", "GET", f"/api/assets/v1/assets?building_id={other}")
    await a("объекты чужого помещения", "GET", f"/api/assets/v1/assets?room_id={room}")
    await a("чужой объект", "GET", f"/api/assets/v1/assets/{asset['id']}")
    await a("история чужого объекта", "GET", f"/api/assets/v1/assets/{asset['id']}/history")
    await a("чужой объект по QR", "GET", f"/api/assets/v1/qr/{asset['qr_code']}")
    await a("люди чужого здания", "GET", f"/api/people/v1/people?building_id={other}")
    await a("люди чужого помещения", "GET", f"/api/people/v1/rooms/{room}/people")
    await a("журнал чужого здания", "GET", f"/api/journal/v1/events?building_id={other}")
    await a("права в чужом здании", "GET", f"/api/access/v1/grants?building_id={other}")
    await a("варианты чужого здания", "GET", f"/api/planning/v1/variants?building_id={other}")
    await a("кабинет чужого здания", "GET", f"/api/views/cabinet?building_id={other}")
    await a("отчёт по чужому зданию", "GET", f"/api/reports/assets.xlsx?building_id={other}")
    await a("лист чужого помещения", "GET", f"/api/reports/rooms/{room}/sheet.pdf")
    await a("изменить чужой объект", "PATCH", f"/api/assets/v1/assets/{asset['id']}", json={"note": "взлом", "version": asset["version"]})
    await a("переместить чужой объект", "POST", f"/api/assets/v1/assets/{asset['id']}/move",
            json={"floor_id": floor, "x": 1, "y": 1, "version": asset["version"]})
    await a("списать чужой объект", "POST", f"/api/assets/v1/assets/{asset['id']}/state", json={"state": "retired", "version": asset["version"]})
    await a("поставить объект в чужое здание", "POST", "/api/assets/v1/assets",
            json={"type_id": asset["type_id"], "floor_id": floor, "x": 1, "y": 1, "inventory_no": f"ВЗЛОМ-{time.time():.0f}"})
    await a("переименовать чужое помещение", "PATCH", f"/api/buildings/v1/rooms/{room}", json={"name": "взлом", "version": 1})
    await a("выдать себе право в чужом здании", "POST", "/api/access/v1/grants",
            json={"person_id": "00000000-0000-0000-0000-000000000000", "role": "admin",
                  "scope": {"kind": "building", "id": other}, "building_id": other})
    await a("вариант в чужом здании", "POST", "/api/planning/v1/variants", json={"building_id": other, "name": "взлом"})
    await a("метка в чужом здании", "POST", f"/api/buildings/v1/buildings/{other}/pins", json={"floor_id": floor, "x": 1, "y": 1, "text": "взлом"})
    await a("удалить чужое здание", "DELETE", f"/api/buildings/v1/buildings/{other}")
    await a("создать здание", "POST", "/api/buildings/v1/buildings", json={"name": "взлом"})
    await a("тип в каталоге", "POST", "/api/catalog/v1/types", json={"name": "взлом"})
    await a("состояние системы", "GET", "/api/system/status")


async def as_viewer(c: httpx.AsyncClient, who: str, own_asset: dict, own_room: str) -> None:
    """In one's own building, what only seeing does not allow."""
    a = lambda *x, **kw: attempt(c, who, *x, **kw)  # noqa: E731
    await a("изменить объект, который только видишь", "PATCH", f"/api/assets/v1/assets/{own_asset['id']}",
            json={"note": "взлом", "version": own_asset["version"]})
    await a("списать объект, который только видишь", "POST", f"/api/assets/v1/assets/{own_asset['id']}/state",
            json={"state": "retired", "version": own_asset["version"]})
    await a("переименовать помещение, которое только видишь", "PATCH", f"/api/buildings/v1/rooms/{own_room}",
            json={"name": "взлом", "version": 1})


def forged_jwt(alg: str) -> str:
    enc = lambda d: base64.urlsafe_b64encode(json.dumps(d).encode()).rstrip(b"=").decode()  # noqa: E731
    now = int(time.time())
    claims = {"sub": "attacker", "iss": f"{BASE}/kc/realms/twin", "aud": "twin", "exp": now + 600, "iat": now,
              "resource_access": {"twin": {"roles": ["twin.admin"]}}, "email": "attacker@example.test"}
    signature = "" if alg == "none" else base64.urlsafe_b64encode(b"x" * 32).rstrip(b"=").decode()
    return f"{enc({'alg': alg, 'typ': 'JWT'})}.{enc(claims)}.{signature}"


async def strangers(other: str) -> None:
    async with httpx.AsyncClient(base_url=BASE) as c:
        for who, headers in (("без входа", {}), ("поддельная cookie", {"Cookie": "twin_session=AAAA.BBBB.CCCC"}),
                             ("поддельный токен", {"Authorization": f"Bearer {forged_jwt('RS256')}"}),
                             ("токен alg=none", {"Authorization": f"Bearer {forged_jwt('none')}"})):
            c.headers.clear()
            c.headers.update(headers)
            await attempt(c, who, "здание", "GET", f"/api/buildings/v1/buildings/{other}/tree")
            await attempt(c, who, "объекты", "GET", f"/api/assets/v1/assets?building_id={other}")
            await attempt(c, who, "люди", "GET", "/api/people/v1/people")
            await attempt(c, who, "создать здание", "POST", "/api/buildings/v1/buildings", json={"name": "взлом"})


def past_the_gateway(project: str, other: str) -> None:
    """Straight at each service from inside the network (from the gateway's container)."""
    script = f"""
import json, urllib.request
out = []
for svc, path in [("buildings", "/api/buildings/v1/buildings/{other}/tree"), ("assets", "/api/assets/v1/assets"),
                  ("people", "/api/people/v1/people"), ("access", "/api/access/v1/grants?building_id={other}"),
                  ("journal", "/api/journal/v1/events"), ("planning", "/api/planning/v1/variants?building_id={other}"),
                  ("import", "/api/import/v1/drawings"), ("catalog", "/api/catalog/v1/types")]:
    for label, token in [("без токена", None), ("поддельный токен", {json.dumps(forged_jwt('RS256'))}),
                         ("токен alg=none", {json.dumps(forged_jwt('none'))})]:
        req = urllib.request.Request(f"http://{{svc}}:8000{{path}}")
        if token:
            req.add_header("Authorization", "Bearer " + token)
        try:
            code = urllib.request.urlopen(req, timeout=10).status
        except urllib.error.HTTPError as e:
            code = e.code
        out.append([svc, label, path, code])
print(json.dumps(out))
"""
    res = subprocess.run(["docker", "exec", "-i", f"{project}-gateway-1", "python", "-"], input=script,
                         capture_output=True, text=True, check=True)
    for svc, label, path, code in json.loads(res.stdout):
        # the catalog is open to everyone signed in, but not to strangers either
        RESULTS.append((code in (401, 403), f"мимо шлюза, {label}", f"{svc}: GET {path.split('?')[0]} → {code}"))


async def main() -> int:
    state = load_state()
    users_of, admins = state["users_of"], set(state["admins"])
    viewer = max(u for u in users_of if u not in admins)            # sees one building
    admin = min(admins)                                             # administers another
    mine = {viewer: users_of[viewer], admin: users_of[admin]}
    other = next(b for b in state["buildings"] if b not in mine.values())
    boss = await sign_in("load-admin", state["admin"]["password"])
    tree = (await boss.get(f"/api/buildings/v1/buildings/{other}/tree")).json()["floors"]
    floor, room = tree[0]["id"], tree[0]["rooms"][0]["id"]
    asset = (await boss.get("/api/assets/v1/assets", params={"building_id": other, "limit": 1})).json()["items"][0]
    own_tree = (await boss.get(f"/api/buildings/v1/buildings/{mine[viewer]}/tree")).json()["floors"]
    own = (await boss.get("/api/assets/v1/assets", params={"building_id": mine[viewer], "limit": 1})).json()["items"][0]
    await boss.aclose()
    for who in (viewer, admin):
        c = await sign_in(who, state["users"][who])
        await as_outsider(c, f"{who} ({'видит' if who == viewer else 'администрирует'} другое здание)", other, room, floor, asset)
        if who == viewer:
            await as_viewer(c, f"{who} в своём здании", own, own_tree[0]["rooms"][0]["id"])
        await c.aclose()
    await strangers(other)
    past_the_gateway(state.get("project", "twin-accept"), other)

    passed = [r for r in RESULTS if r[0]]
    failed = [r for r in RESULTS if not r[0]]
    for _, who, what in failed:
        print(f"ПРОШЛО  {who}: {what}")
    print(f"попыток {len(RESULTS)}, отклонено {len(passed)}, прошло {len(failed)}")
    print("ПРИНЯТО" if not failed else "НЕ ПРИНЯТО")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
