"""Requests of the people service, under /api/people/v1 (contracts/openapi/people.v1.json).

The directory is one for the organisation. It is kept by administrators of
buildings (anyone who is the administrator of at least one) and by the system
administrator; who sees whom by rooms arrives with assignments to rooms (stage 6).
Name and position of a person are open to anyone signed in, by id — to show who
answers for an object and who made a change.
"""
from __future__ import annotations

import io
from uuid import UUID

from fastapi import APIRouter, Depends, File, Request, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import func, insert, select, update
from twin_chassis import Identity, Problem, current_identity, emit, uuid7
from twin_chassis.access import decide

from .models import persons

router = APIRouter()
MAX_IMPORT_ROWS = 10000
# something@domain.zone — internal domains (.local, .corp) are as good as public ones
EMAIL = r"^[^@\s]+@[^@\s]+\.[^@\s]+$"


def person_json(p, *, brief: bool = False) -> dict:
    out = {"id": str(p.id), "full_name": p.full_name, "position": p.position}
    if not brief:
        out.update(department=p.department, email=p.email, phone=p.phone, linked=p.user_sub is not None,
                   status=p.status, version=p.version)
    return out


async def _keeper(request: Request, me: Identity) -> None:
    """The directory is kept by administrators of buildings and the system administrator."""
    if me.is_admin or me.kind == "service":
        return
    scopes = await request.app.state.rights.scopes(request, me, reading=False)
    if not any(g["role"] == "admin" and g["scope"]["kind"] == "building" for g in scopes.get("grants") or []):
        raise Problem(403, "forbidden", "Справочник людей ведут администраторы зданий")


async def _person(session, person_id: UUID, *, lock: bool = False):
    q = select(persons).where(persons.c.id == person_id)
    p = (await session.execute(q.with_for_update() if lock else q)).one_or_none()
    if p is None:
        raise Problem(404, "not_found", "Человек не найден")
    return p


def _stale(p, version: int) -> None:
    if p.version != version:
        raise Problem(409, "version_conflict", "Запись изменил кто-то другой — обновите страницу",
                      extra={"current": person_json(p)})


# ----------------------------------------------------------------------- me ---
@router.get("/me", tags=["me"], summary="Запись справочника для того, кто вошёл; при первом входе связывает по почте")
async def me_(request: Request, me: Identity = Depends(current_identity)) -> dict:
    """`person` is null when the account is not in the directory; `reason` says why:
    `unknown` — nobody with this e-mail, `ambiguous` — several people share it (an
    administrator links by hand), `left` — the person has left."""
    async with request.app.state.db.begin() as session:
        p = (await session.execute(select(persons).where(persons.c.user_sub == me.sub))).one_or_none()
        if p is not None:
            return {"person": person_json(p) if p.status == "active" else None,
                    "reason": None if p.status == "active" else "left"}
        email = (me.email or "").strip().lower()
        found = [] if not email else (await session.execute(select(persons).where(
            persons.c.email == email, persons.c.status == "active", persons.c.user_sub.is_(None))
            .with_for_update())).all()
        if len(found) != 1:
            return {"person": None, "reason": "ambiguous" if len(found) > 1 else "unknown"}
        p = (await session.execute(update(persons).where(persons.c.id == found[0].id).values(
            user_sub=me.sub, version=persons.c.version + 1, updated_at=func.now()).returning(persons))).one()
        await emit(session, "people.person.linked", ("person", p.id), {"person_id": str(p.id), "user_sub": me.sub},
                   person_id=p.id)
        return {"person": person_json(p), "reason": None}


# ------------------------------------------------------------------- people ---
class PersonIn(BaseModel):
    full_name: str = Field(min_length=1, max_length=200)
    position: str = Field(default="", max_length=200)
    department: str = Field(default="", max_length=200)
    email: str | None = Field(default=None, max_length=200, pattern=EMAIL)
    phone: str | None = Field(default=None, max_length=40)


class PersonPatch(BaseModel):
    full_name: str | None = Field(default=None, min_length=1, max_length=200)
    position: str | None = Field(default=None, max_length=200)
    department: str | None = Field(default=None, max_length=200)
    email: str | None = Field(default=None, max_length=200, pattern=EMAIL)
    phone: str | None = Field(default=None, max_length=40)
    version: int


def _clean(values: dict) -> dict:
    out = {k: (v.strip() if isinstance(v, str) else v) for k, v in values.items()}
    if out.get("email"):
        out["email"] = out["email"].lower()
    return out


@router.get("/people", tags=["people"], summary="Справочник людей")
async def list_people(request: Request, q: str | None = None, department: str | None = None, status: str = "active",
                      building_id: UUID | None = None, room_id: UUID | None = None, limit: int = 500,
                      me: Identity = Depends(current_identity)) -> dict:
    """The keepers of the directory see everyone. Narrowed by `building_id` or `room_id` to the people assigned
    there, it shows who sits where — that only in rooms the asker may see (the system administrator sees all);
    anyone else sees, with `building_id`, the people of the rooms they have a right to."""
    try:
        await _keeper(request, me)
        keeper = me.is_admin or me.kind == "service" or (building_id is None and room_id is None)
    except Problem:
        keeper = False
        if building_id is None:
            raise Problem(403, "forbidden", "Справочник целиком ведут администраторы зданий; укажите здание") from None
    where = [] if status == "all" else [persons.c.status == status]
    if building_id is not None or room_id is not None:
        from .assignments import chain
        from .models import assignments
        if room_id is not None:
            where.append(persons.c.id.in_(select(assignments.c.person_id).where(assignments.c.room_id == room_id)))
        if building_id is not None:
            where.append(persons.c.id.in_(
                select(assignments.c.person_id).where(assignments.c.building_id == building_id)))
    if department:
        where.append(persons.c.department == department)
    if q and q.strip():
        like = "%" + q.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
        where.append(persons.c.full_name.ilike(like) | func.coalesce(persons.c.email, "").ilike(like))
    async with request.app.state.db.begin() as session:
        rows = (await session.execute(select(persons).where(*where).order_by(persons.c.full_name, persons.c.id)
                                      .limit(max(1, min(limit, 500))))).all()
        if not keeper:                                   # only those in rooms the asker may see
            scopes = await request.app.state.rights.scopes(request, me)
            seen = (await session.execute(select(assignments).where(
                assignments.c.person_id.in_([p.id for p in rows] or [uuid7()])))).all()
            visible = {a.person_id for a in seen if decide(scopes, "view", chain(a))}
            rows = [p for p in rows if p.id in visible]
    return {"items": [person_json(p) for p in rows], "next": None}


@router.post("/people", status_code=201, tags=["people"], summary="Добавить человека")
async def create_person(body: PersonIn, request: Request, me: Identity = Depends(current_identity)) -> dict:
    await _keeper(request, me)
    async with request.app.state.db.begin() as session:
        p = (await session.execute(insert(persons).values(id=uuid7(), created_by=me.sub, **_clean(body.model_dump()))
                                   .returning(persons))).one()
        await emit(session, "people.person.created", ("person", p.id), {"person_id": str(p.id)}, person_id=p.id)
    return person_json(p)


class Ids(BaseModel):
    ids: list[UUID] = Field(max_length=1000)


@router.post("/people/batch-get", tags=["people"], summary="Имя и должность людей по идентификаторам")
async def batch_people(body: Ids, request: Request, _: Identity = Depends(current_identity)) -> dict:
    async with request.app.state.db.begin() as session:
        rows = (await session.execute(select(persons).where(persons.c.id.in_(body.ids or [uuid7()])))).all()
    return {"items": [person_json(p, brief=True) for p in rows]}


@router.get("/people/{person_id}", tags=["people"], summary="Человек")
async def get_person(person_id: UUID, request: Request, me: Identity = Depends(current_identity)) -> dict:
    """With contacts — to the keepers of the directory, to the person themself and to whoever sees a room
    the person is assigned to; to the rest the person is not there (404)."""
    from .assignments import chain
    from .models import assignments
    async with request.app.state.db.begin() as session:
        p = await _person(session, person_id)
        try:
            await _keeper(request, me)
            return person_json(p)
        except Problem:
            pass
        scopes = await request.app.state.rights.scopes(request, me)
        rows = (await session.execute(select(assignments).where(assignments.c.person_id == person_id))).all()
        if scopes.get("person_id") == str(person_id) or any(decide(scopes, "view", chain(a)) for a in rows):
            return person_json(p)
    raise Problem(404, "not_found", "Человек не найден")


@router.patch("/people/{person_id}", tags=["people"], summary="Изменить запись")
async def patch_person(person_id: UUID, body: PersonPatch, request: Request,
                       me: Identity = Depends(current_identity)) -> dict:
    await _keeper(request, me)
    values = _clean(body.model_dump(exclude={"version"}, exclude_unset=True))
    async with request.app.state.db.begin() as session:
        p = await _person(session, person_id, lock=True)
        _stale(p, body.version)
        if not values:
            return person_json(p)
        p = (await session.execute(update(persons).where(persons.c.id == person_id).values(
            **values, version=persons.c.version + 1, updated_at=func.now()).returning(persons))).one()
        await emit(session, "people.person.changed", ("person", p.id),
                   {"person_id": str(p.id), "fields": sorted(values)}, person_id=p.id)
    return person_json(p)


@router.post("/people/{person_id}/leave", tags=["people"], summary="Человек выбыл: его права отзываются")
async def leave(person_id: UUID, request: Request, me: Identity = Depends(current_identity)) -> dict:
    await _keeper(request, me)
    async with request.app.state.db.begin() as session:
        p = await _person(session, person_id, lock=True)
        if p.status == "left":
            return person_json(p)
        p = (await session.execute(update(persons).where(persons.c.id == person_id).values(
            status="left", left_at=func.now(), version=persons.c.version + 1, updated_at=func.now())
            .returning(persons))).one()
        await emit(session, "people.person.left", ("person", p.id), {"person_id": str(p.id)}, person_id=p.id)
    return person_json(p)


class Link(BaseModel):
    user_sub: str = Field(min_length=1, max_length=200, description="Идентификатор учётной записи в Keycloak")


@router.put("/people/{person_id}/link", tags=["people"], summary="Связать человека с учётной записью вручную")
async def link(person_id: UUID, body: Link, request: Request, me: Identity = Depends(current_identity)) -> dict:
    await _keeper(request, me)
    async with request.app.state.db.begin() as session:
        p = await _person(session, person_id, lock=True)
        if p.status != "active":
            raise Problem(409, "left", "Человек выбыл — связать его с учётной записью нельзя")
        taken = (await session.execute(select(persons.c.full_name).where(
            persons.c.user_sub == body.user_sub, persons.c.id != person_id))).scalar()
        if taken:
            raise Problem(409, "account_taken", f"Эта учётная запись уже связана с человеком «{taken}»")
        p = (await session.execute(update(persons).where(persons.c.id == person_id).values(
            user_sub=body.user_sub, version=persons.c.version + 1, updated_at=func.now()).returning(persons))).one()
        await emit(session, "people.person.linked", ("person", p.id),
                   {"person_id": str(p.id), "user_sub": body.user_sub}, person_id=p.id)
    return person_json(p)


@router.delete("/people/{person_id}/link", tags=["people"], summary="Разорвать связь с учётной записью")
async def unlink(person_id: UUID, request: Request, me: Identity = Depends(current_identity)) -> dict:
    await _keeper(request, me)
    async with request.app.state.db.begin() as session:
        p = await _person(session, person_id, lock=True)
        if p.user_sub is None:
            return person_json(p)
        sub = p.user_sub
        p = (await session.execute(update(persons).where(persons.c.id == person_id).values(
            user_sub=None, version=persons.c.version + 1, updated_at=func.now()).returning(persons))).one()
        await emit(session, "people.person.unlinked", ("person", p.id), {"person_id": str(p.id), "user_sub": sub},
                   person_id=p.id)
    return person_json(p)


# ------------------------------------------------------------------- import ---
COLUMNS = {"фио": "full_name", "должность": "position", "подразделение": "department",
           "почта": "email", "email": "email", "e-mail": "email", "телефон": "phone"}


@router.post("/people/import", tags=["people"],
             summary="Загрузить людей из Excel: столбцы ФИО, Должность, Подразделение, Почта, Телефон")
async def import_people(request: Request, file: UploadFile = File(...),
                        me: Identity = Depends(current_identity)) -> dict:
    """A person is recognised by the e-mail, without one — by name and department, so the
    same file loaded again creates nobody. Rows with errors are reported and skipped."""
    from openpyxl import load_workbook

    await _keeper(request, me)
    try:
        sheet = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as ex:
        raise Problem(422, "invalid", "Не удалось прочитать файл: нужен Excel (.xlsx)",
                      errors=[{"field": "file", "message": "не .xlsx"}]) from ex
    if not rows:
        raise Problem(422, "invalid", "Файл пуст", errors=[{"field": "file", "message": "нет строк"}])
    header = [COLUMNS.get(str(c or "").strip().lower()) for c in rows[0]]
    if "full_name" not in header:
        raise Problem(422, "invalid", "В первой строке нужен столбец «ФИО»",
                      errors=[{"field": "file", "message": "нет столбца ФИО"}])
    if len(rows) - 1 > MAX_IMPORT_ROWS:
        raise Problem(422, "invalid", f"Строк больше {MAX_IMPORT_ROWS}")
    report = {"created": 0, "updated": 0, "unchanged": 0, "errors": []}
    async with request.app.state.db.begin() as session:
        for n, row in enumerate(rows[1:], start=2):
            raw = {k: (str(v).strip() if v is not None else "") for k, v in zip(header, row, strict=False) if k}
            if not any(raw.values()):
                continue
            try:
                values = _clean(PersonIn(**{k: (v or None) if k in ("email", "phone") else v
                                            for k, v in raw.items()}).model_dump())
            except ValueError as ex:
                first = ex.errors()[0] if hasattr(ex, "errors") else {"loc": ["?"], "msg": str(ex)}
                report["errors"].append({"row": n, "message": f"{'.'.join(map(str, first['loc']))}: {first['msg']}"})
                continue
            same = (persons.c.email == values["email"]) if values["email"] else (
                (persons.c.full_name == values["full_name"]) & (persons.c.department == values["department"])
                & persons.c.email.is_(None))
            old = (await session.execute(select(persons).where(same, persons.c.status == "active"))).first()
            if old is None:
                p = (await session.execute(insert(persons).values(id=uuid7(), created_by=me.sub, **values)
                                           .returning(persons.c.id))).scalar_one()
                await emit(session, "people.person.created", ("person", p), {"person_id": str(p)}, person_id=p)
                report["created"] += 1
                continue
            changed = sorted(k for k, v in values.items() if (getattr(old, k) or None) != (v or None))
            if not changed:
                report["unchanged"] += 1
                continue
            await session.execute(update(persons).where(persons.c.id == old.id).values(
                **values, version=persons.c.version + 1, updated_at=func.now()))
            await emit(session, "people.person.changed", ("person", old.id),
                       {"person_id": str(old.id), "fields": changed}, person_id=old.id)
            report["updated"] += 1
    return report
